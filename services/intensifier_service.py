from __future__ import annotations

from dataclasses import dataclass
from pathlib import Path
import re

import numpy as np
import openpyxl

from PLIFMix.algorithms.intensifier import apply_correction_matrix, fit_linear_gain_matrix
from PLIFMix.core.image_io import extract_plif_signal, list_image_files, read_image, save_float_image
from PLIFMix.core.paths import path_is_set
from PLIFMix.domain.settings import GainCorrectionSettings

GAIN_SCALE = 100


@dataclass
class CorrectionResult:
    processed_count: int
    outputs: list[dict[str, Path]]
    records: list["CorrectionRecord"]
    unmatched_paths: tuple[Path, ...] = ()


@dataclass
class CorrectionRecord:
    image_path: Path
    matrix_path: Path
    gain: float | None
    outputs: dict[str, Path]


@dataclass
class MatrixFitResult:
    image_count: int
    outputs: dict[str, Path]


class IntensifierCorrectionService:
    def apply_folder(self, settings: GainCorrectionSettings) -> CorrectionResult:
        image_folder = Path(settings.image_folder)
        matrix_path = Path(settings.matrix_path)
        if not image_folder.exists():
            raise FileNotFoundError(f"Image folder does not exist: {image_folder}")
        if not matrix_path.exists():
            raise FileNotFoundError(f"Correction matrix does not exist: {matrix_path}")
        files = list_image_files(image_folder, (".npy",))
        return self.apply_paths(settings, files)

    def apply_paths(self, settings: GainCorrectionSettings, paths: list[Path]) -> CorrectionResult:
        if not paths:
            raise ValueError("No image files selected.")
        paths = [Path(path) for path in paths]
        _ensure_unique_output_names(paths)
        matrix_path = Path(settings.matrix_path)
        if not matrix_path.exists():
            raise FileNotFoundError(f"Correction matrix does not exist: {matrix_path}")
        matrix = np.load(matrix_path)
        output_folder = Path(settings.output_folder) if path_is_set(settings.output_folder) else Path(paths[0]).parent.parent / "intensifier_corrected"
        records: list[CorrectionRecord] = []
        for path in paths:
            outputs = apply_correction_matrix(path, matrix, output_folder)
            records.append(CorrectionRecord(path, matrix_path, None, outputs))
        return CorrectionResult(len(records), [record.outputs for record in records], records)

    def apply_by_excel(self, settings: GainCorrectionSettings) -> CorrectionResult:
        image_folder = Path(settings.image_folder)
        excel_path = Path(settings.excel_path)
        matrix_folder = Path(settings.matrix_folder)
        if settings.experiment_column < 1 or settings.gain_column < 1:
            raise ValueError("Excel column numbers must be positive integers.")
        if not image_folder.exists():
            raise FileNotFoundError(f"Image folder does not exist: {image_folder}")
        if not excel_path.exists():
            raise FileNotFoundError(f"Excel file does not exist: {excel_path}")
        if not matrix_folder.exists():
            raise FileNotFoundError(f"Matrix folder does not exist: {matrix_folder}")
        workbook = openpyxl.load_workbook(excel_path, read_only=True, data_only=True)
        try:
            sheet = workbook.active
            gain_by_experiment: dict[int, int] = {}
            for row in range(1, sheet.max_row + 1):
                exp_value = sheet.cell(row=row, column=settings.experiment_column).value
                gain_value = sheet.cell(row=row, column=settings.gain_column).value
                if exp_value is None or gain_value is None:
                    continue
                try:
                    gain_by_experiment[int(exp_value)] = int(float(gain_value) * GAIN_SCALE)
                except (TypeError, ValueError):
                    continue
        finally:
            workbook.close()

        output_folder = Path(settings.output_folder) if path_is_set(settings.output_folder) else image_folder.parent / "intensifier_corrected"
        records: list[CorrectionRecord] = []
        image_paths = sorted(path for path in image_folder.rglob("*.npy") if path.is_file())
        unmatched_paths: list[Path] = []
        matched_paths: list[tuple[Path, int, int]] = []
        for image_path in image_paths:
            exp_no = _experiment_number(image_path.stem)
            if exp_no is None or exp_no not in gain_by_experiment:
                unmatched_paths.append(image_path)
            else:
                matched_paths.append((image_path, exp_no, gain_by_experiment[exp_no]))
        _ensure_unique_output_names([path for path, _exp_no, _gain in matched_paths])

        matrix_cache: dict[int, np.ndarray] = {}
        for image_path, exp_no, gain in matched_paths:
            matrix_path = matrix_folder / f"Gain{gain}.npy"
            if not matrix_path.exists():
                raise FileNotFoundError(f"Experiment {exp_no}: matrix not found: {matrix_path}")
            matrix = matrix_cache.get(gain)
            if matrix is None:
                matrix = np.load(matrix_path)
                matrix_cache[gain] = matrix
            outputs = apply_correction_matrix(image_path, matrix, output_folder)
            records.append(CorrectionRecord(image_path, matrix_path, gain, outputs))
        return CorrectionResult(
            len(records),
            [record.outputs for record in records],
            records,
            tuple(unmatched_paths),
        )

    def fit_matrix_from_folder(self, folder: str | Path, output_folder: str | Path | None = None) -> MatrixFitResult:
        folder = Path(folder)
        files = list_image_files(folder, (".png", ".tif", ".tiff", ".bmp", ".npy"))
        if not files:
            raise ValueError(f"No images found for matrix fitting: {folder}")
        x_values = np.array([_parse_numeric_stem(path) for path in files], dtype=np.float64)
        # v1 3.3 fits one scalar image per exposure.  Keep .npy inputs unchanged
        # and extract the same green PLIF plane when calibration inputs are color images.
        stack = np.dstack([extract_plif_signal(read_image(path)) for path in files])
        _, correction = fit_linear_gain_matrix(stack, x_values)
        out = Path(output_folder) if output_folder is not None and path_is_set(output_folder) else folder.parent
        outputs = save_float_image(correction, out, folder.name)
        return MatrixFitResult(image_count=len(files), outputs=outputs)


def _parse_numeric_stem(path: Path) -> int:
    stem = path.stem
    if stem.isdigit():
        return int(stem)
    digits = "".join(ch for ch in stem if ch.isdigit())
    if not digits:
        raise ValueError(f"Cannot parse numeric exposure/gain from file name: {path.name}")
    return int(digits)


def _experiment_number(stem: str) -> int | None:
    match = re.search(r"实验编号\s*(\d+)", stem)
    if match:
        return int(match.group(1))
    match = re.search(r"(\d+)", stem)
    return int(match.group(1)) if match else None


def _ensure_unique_output_names(paths: list[Path]) -> None:
    by_stem: dict[str, Path] = {}
    for path in paths:
        previous = by_stem.get(path.stem)
        if previous is not None and previous.resolve() != path.resolve():
            raise FileExistsError(
                f"输出文件名冲突: {previous} 与 {path} 都会保存为 {path.stem}.npy"
            )
        by_stem[path.stem] = path
