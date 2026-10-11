from __future__ import annotations

from pathlib import Path

import openpyxl
import numpy as np
from matplotlib.backends.backend_agg import FigureCanvasAgg
from matplotlib.figure import Figure

from PLIFMix.core.paths import path_is_set
from PLIFMix.core.safe_io import atomic_save_workbook, atomic_write, excel_write_session
from PLIFMix.domain.settings import DistributionSettings
from PLIFMix.services.distribution_results import DistributionFitParameters, DistributionFitResult


def export_distribution_result(settings: DistributionSettings, result: DistributionFitResult) -> Path:
    output = (
        Path(settings.output_folder)
        if path_is_set(settings.output_folder)
        else Path(settings.image_folder).parent / "distribution_results"
    )
    output.mkdir(parents=True, exist_ok=True)
    name = settings.save_name or f"F{settings.f_exp_no}_FO{settings.fo_exp_no}_O{settings.o_exp_no}"

    workbook = openpyxl.Workbook()
    sheet = workbook.active
    sheet.title = "fit_parameters"
    headers = [
        "center_x",
        "center_y",
        "rotate_angle",
        "center_a",
        "center_b",
        "center_r2",
        "hori_a",
        "hori_b",
        "hori_c",
        "hori_r2",
        "vert_a",
        "vert_b",
        "vert_c",
        "vert_r2",
        "student_scale_x",
        "student_scale_y",
        "student_exponent",
        "student_r2",
        "compact_scale_x",
        "compact_scale_y",
        "compact_exponent",
        "compact_cutoff_radius",
        "compact_r2",
    ]
    sheet.append(["image", *headers])
    for fit in result.numeric_results:
        sheet.append([fit.image_name, *_fit_parameter_values(fit)])
    workbook_path = output / f"{name}-fit-results.xlsx"
    with excel_write_session(workbook_path):
        atomic_save_workbook(workbook, workbook_path)

    for curve_name, (x_values, values, predicted) in result.curve_data.items():
        figure = Figure(figsize=(8, 5), dpi=150)
        axis = figure.add_subplot(111)
        axis.scatter(x_values, values, s=6, label="Origin")
        axis.plot(x_values, predicted, color="red", label="Fit")
        axis.set_xlabel("dimensionless coordinate")
        axis.set_ylabel("normalized intensity")
        axis.grid(True, alpha=0.3)
        axis.legend()
        FigureCanvasAgg(figure)
        atomic_write(
            output / f"{name}-{curve_name}.png",
            lambda temporary: figure.savefig(temporary, format="png", bbox_inches="tight"),
            "save fit curve image",
        )
    return output


def write_fit_results_to_excel(settings: DistributionSettings, result: DistributionFitResult) -> Path:
    """Write the current F/O/FO fit values into the existing measurement workbook."""
    target = Path(settings.fit_result_path)
    if not target.is_file():
        raise ValueError(f"拟合结果 Excel 文件不存在: {target}")
    if len(result.numeric_results) < 3:
        raise ValueError("当前拟合结果不完整，至少需要 F、O、FO 三行数据。")

    fits_by_name = {fit.image_name: fit for fit in result.numeric_results}
    try:
        f_fit = fits_by_name["f_img"]
        o_fit = fits_by_name["o_img"]
        fo_fit = fits_by_name["fo_img"]
    except KeyError as exc:
        raise ValueError("当前拟合结果缺少 F、O 或 FO 图像数据。") from exc

    try:
        with excel_write_session(target):
            workbook = openpyxl.load_workbook(target, data_only=False)
            workbook_open = True
            try:
                worksheet = workbook["测量数据"] if "测量数据" in workbook.sheetnames else workbook.active
                expected_headers = (
                    "中心x",
                    "中心y",
                    "旋转角度",
                    "上界",
                    "下界",
                    "中心线a",
                    "中心线b",
                    "R2",
                    "横向a",
                    "横向b",
                    "横向c",
                    "R2",
                    "纵向a",
                    "纵向b",
                    "纵向c",
                    "R2",
                    "sx",
                    "sy",
                    "m",
                    "R2",
                    "sx",
                    "sy",
                    "m",
                    "R",
                    "R2",
                )
                actual_headers = tuple(worksheet.cell(row=1, column=column).value for column in range(20, 45))
                if actual_headers != expected_headers:
                    raise ValueError("目标 Excel 的 T:AR 表头与拟合结果模板不一致。")

                requested = (
                    (
                        _experiment_number_key(settings.f_exp_no),
                        f_fit,
                        settings.q_center_up_f,
                        settings.q_center_down_f,
                    ),
                    (
                        _experiment_number_key(settings.o_exp_no),
                        o_fit,
                        settings.q_center_up_o,
                        settings.q_center_down_o,
                    ),
                    (
                        _experiment_number_key(settings.fo_exp_no),
                        fo_fit,
                        settings.q_center_up_fo,
                        settings.q_center_down_fo,
                    ),
                )
                rows_by_number: dict[str, list[int]] = {}
                for row_index in range(1, worksheet.max_row + 1):
                    key = _experiment_number_key(worksheet.cell(row=row_index, column=1).value)
                    if key:
                        rows_by_number.setdefault(key, []).append(row_index)

                target_rows: list[tuple[int, DistributionFitParameters, int, int]] = []
                for number, fit, upper, lower in requested:
                    if not number:
                        raise ValueError("F、O、FO 实验编号不能为空。")
                    matches = rows_by_number.get(number, [])
                    if not matches:
                        raise ValueError(f"Excel A 列中找不到实验编号 {number}。")
                    if len(matches) > 1:
                        raise ValueError(f"Excel A 列中实验编号 {number} 出现多次，无法确定写入行。")
                    target_rows.append((matches[0], fit, upper, lower))

                for row_index, fit, upper, lower in target_rows:
                    values = [
                        fit.center_x,
                        fit.center_y,
                        fit.rotate_angle,
                        upper,
                        lower,
                        fit.center_a,
                        fit.center_b,
                        fit.center_r_squared,
                        fit.horizontal_a,
                        fit.horizontal_b,
                        fit.horizontal_c,
                        fit.horizontal_r_squared,
                        fit.vertical_a,
                        fit.vertical_b,
                        fit.vertical_c,
                        fit.vertical_r_squared,
                        fit.student_scale_x,
                        fit.student_scale_y,
                        fit.student_exponent,
                        fit.student_r_squared,
                        fit.compact_scale_x,
                        fit.compact_scale_y,
                        fit.compact_exponent,
                        fit.compact_cutoff_radius,
                        fit.compact_r_squared,
                    ]
                    for column, value in enumerate(values, start=20):
                        worksheet.cell(row=row_index, column=column).value = value
                atomic_save_workbook(workbook, target)
                workbook_open = False
            finally:
                if workbook_open:
                    workbook.close()
    except PermissionError as exc:
        raise ValueError("拟合结果 Excel 可能正在被其他程序占用，请关闭后重试。") from exc
    return target


def _fit_parameter_values(fit: DistributionFitParameters) -> list[float | None]:
    return [
        fit.center_x,
        fit.center_y,
        fit.rotate_angle,
        fit.center_a,
        fit.center_b,
        fit.center_r_squared,
        fit.horizontal_a,
        fit.horizontal_b,
        fit.horizontal_c,
        fit.horizontal_r_squared,
        fit.vertical_a,
        fit.vertical_b,
        fit.vertical_c,
        fit.vertical_r_squared,
        fit.student_scale_x,
        fit.student_scale_y,
        fit.student_exponent,
        fit.student_r_squared,
        fit.compact_scale_x,
        fit.compact_scale_y,
        fit.compact_exponent,
        fit.compact_cutoff_radius,
        fit.compact_r_squared,
    ]


def _experiment_number_key(value) -> str:
    if value is None:
        return ""
    if isinstance(value, (int, np.integer)):
        return str(int(value))
    if isinstance(value, (float, np.floating)) and np.isfinite(value) and float(value).is_integer():
        return str(int(value))
    text = str(value).strip()
    if text.startswith("实验编号"):
        text = text[len("实验编号") :].strip()
    try:
        numeric = float(text)
    except ValueError:
        return text
    return str(int(numeric)) if numeric.is_integer() else text


__all__ = ["export_distribution_result", "write_fit_results_to_excel"]
