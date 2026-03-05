import openpyxl
def write_excel_func(x, y, filepath):
    wb = openpyxl.Workbook()
    ws = wb.active
    for i in range(len(x)):
        ws.cell(row=i+1, column=1, value=x[i])
        ws.cell(row=i+1, column=2, value=y[i])
    wb.save(filepath)