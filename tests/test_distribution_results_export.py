from __future__ import annotations

from dataclasses import replace
from pathlib import Path
from tempfile import TemporaryDirectory
import unittest

import numpy as np
import openpyxl

from PLIFv2.core.result_snapshot import ResultSnapshot
from PLIFv2.domain.settings import DistributionSettings
from PLIFv2.services.distribution_presentation import (
    format_center_line_fit,
    format_profile_fit,
    format_surface_fits,
)
from PLIFv2.services.distribution_service import DistributionFitService
from PLIFv2.services.distribution_results import DistributionFitParameters, DistributionFitResult
from PLIFv2.services.distribution_workers import SurfaceFitParameters


def _fit(image_name: str, start: float) -> DistributionFitParameters:
    return DistributionFitParameters(
        image_name=image_name,
        center_x=start,
        center_y=start + 1,
        rotate_angle=start + 2,
        center_a=start + 3,
        center_b=start + 4,
        center_r_squared=start + 5,
        horizontal_a=start + 6,
        horizontal_b=start + 7,
        horizontal_c=start + 8,
        horizontal_r_squared=start + 9,
        vertical_a=start + 10,
        vertical_b=start + 11,
        vertical_c=start + 12,
        vertical_r_squared=start + 13,
        student_scale_x=start + 14,
        student_scale_y=start + 15,
        student_exponent=start + 16,
        student_r_squared=start + 17,
        compact_scale_x=start + 18,
        compact_scale_y=start + 19,
        compact_exponent=start + 20,
        compact_cutoff_radius=start + 21,
        compact_r_squared=start + 22,
    )


def _result() -> DistributionFitResult:
    return DistributionFitResult(
        fit_text="fit",
        surface_text="surface",
        images=[],
        numeric_results=(
            _fit("f_img", 10.0),
            _fit("o_img", 40.0),
            _fit("fo_img", 70.0),
        ),
        curve_data={
            "横向分布_f_img": (
                np.array([-1.0, 0.0, 1.0]),
                np.array([0.1, 0.4, 0.1]),
                np.array([0.2, 0.3, 0.2]),
            )
        },
        warnings=(),
        summary_rows=(),
    )


class DistributionResultExportTests(unittest.TestCase):
    def test_presenter_keeps_legacy_fit_text_layout(self) -> None:
        center = format_center_line_fit(1.23456, 2.34567, 0.98765)
        horizontal = format_profile_fit("horizontal", 0.123456, 2.34567, 3.45678, 0.87654)
        vertical = format_profile_fit("vertical", 0.223456, 2.44567, 3.55678, 0.77654)
        surface = format_surface_fits(
            "f_img",
            SurfaceFitParameters(1.1, 1.2, 1.3, None, 0.91),
            SurfaceFitParameters(1.4, 1.5, 1.6, 4.0, 0.92),
            None,
            None,
            True,
        )

        self.assertEqual(center.text, "x=a(√(y²+b²)-b); a=1.2346; b=2.3457; R²=0.9877")
        self.assertEqual(
            horizontal.text,
            "y=1234.5600E-4/(x^2*2.3457^2+1)^3.4568; R²=0.8765",
        )
        self.assertEqual(
            vertical.equation,
            "p=a/(y²b²+1)^c",
        )
        self.assertEqual(
            surface.text,
            "f_img纯椭圆 Student：sₓ=1.1000; sᵧ=1.2000; m=1.3000; R²=0.9100\n"
            "f_img紧支撑 Student：sₓ=1.4000; sᵧ=1.5000; m=1.6000; R=4.0000; R²=0.9200\n",
        )

    def test_named_numeric_results_keep_legacy_fit_rows_read_view(self) -> None:
        result = _result()

        self.assertEqual(result.numeric_results[1].horizontal_a, 46.0)
        self.assertEqual(result.fit_rows, [item.as_fit_row() for item in result.numeric_results])
        self.assertEqual(result.fit_rows[2][14:], [84.0, 85.0, 86.0, 87.0, 88.0, 89.0, 90.0, 91.0, 92.0])

    def test_provenance_can_be_replaced_without_mutating_cached_result(self) -> None:
        result = _result()
        snapshot = ResultSnapshot.capture(DistributionSettings(), ())

        replaced = replace(result, snapshot=snapshot)

        self.assertIsNone(result.snapshot)
        self.assertIs(replaced.snapshot, snapshot)
        self.assertIs(replaced.numeric_results, result.numeric_results)

    def test_service_export_writes_named_values_and_curve_plot(self) -> None:
        result = _result()
        with TemporaryDirectory() as folder:
            root = Path(folder)
            settings = DistributionSettings(
                image_folder=root / "images",
                output_folder=root / "output",
                save_name="named",
                f_exp_no="1",
                o_exp_no="2",
                fo_exp_no="3",
            )
            service = DistributionFitService()
            output = service.export(settings, result)

            workbook_path = output / "named-fit-results.xlsx"
            self.assertEqual(output, root / "output")
            self.assertTrue(workbook_path.is_file())
            self.assertTrue((output / "named-横向分布_f_img.png").is_file())
            workbook = openpyxl.load_workbook(workbook_path, data_only=False)
            try:
                worksheet = workbook["fit_parameters"]
                self.assertEqual(
                    [worksheet.cell(row=2, column=column).value for column in range(1, 25)],
                    ["f_img", *result.numeric_results[0].as_fit_row()],
                )
            finally:
                workbook.close()

    def test_service_measurement_export_uses_named_fields_in_matching_rows(self) -> None:
        result = _result()
        headers = (
            "中心x",
            "中心y",
            "旋转角度",
            "上界",
            "下界",
            "中心线a",
            "中心线b",
            "R2",
            "横向a",
            "横向b",
            "横向c",
            "R2",
            "纵向a",
            "纵向b",
            "纵向c",
            "R2",
            "sx",
            "sy",
            "m",
            "R2",
            "sx",
            "sy",
            "m",
            "R",
            "R2",
        )
        with TemporaryDirectory() as folder:
            root = Path(folder)
            workbook_path = root / "measurement.xlsx"
            workbook = openpyxl.Workbook()
            worksheet = workbook.active
            worksheet.title = "测量数据"
            for column, value in enumerate(headers, start=20):
                worksheet.cell(row=1, column=column).value = value
            for row, number in enumerate((3, 1, 2), start=2):
                worksheet.cell(row=row, column=1).value = number
            workbook.save(workbook_path)
            workbook.close()

            settings = DistributionSettings(
                image_folder=root / "images",
                fit_result_path=workbook_path,
                f_exp_no="1",
                o_exp_no="2",
                fo_exp_no="3",
                q_center_up_f=11,
                q_center_down_f=12,
                q_center_up_o=21,
                q_center_down_o=22,
                q_center_up_fo=31,
                q_center_down_fo=32,
            )
            DistributionFitService().write_fit_results_to_excel(settings, result)

            workbook = openpyxl.load_workbook(workbook_path, data_only=False)
            try:
                worksheet = workbook["测量数据"]
                fits = {
                    1: (result.numeric_results[0], 11, 12),
                    2: (result.numeric_results[1], 21, 22),
                    3: (result.numeric_results[2], 31, 32),
                }
                for row, number in enumerate((3, 1, 2), start=2):
                    fit, upper, lower = fits[number]
                    values = [
                        fit.center_x,
                        fit.center_y,
                        fit.rotate_angle,
                        upper,
                        lower,
                        fit.center_a,
                        fit.center_b,
                        fit.center_r_squared,
                        fit.horizontal_a,
                        fit.horizontal_b,
                        fit.horizontal_c,
                        fit.horizontal_r_squared,
                        fit.vertical_a,
                        fit.vertical_b,
                        fit.vertical_c,
                        fit.vertical_r_squared,
                        fit.student_scale_x,
                        fit.student_scale_y,
                        fit.student_exponent,
                        fit.student_r_squared,
                        fit.compact_scale_x,
                        fit.compact_scale_y,
                        fit.compact_exponent,
                        fit.compact_cutoff_radius,
                        fit.compact_r_squared,
                    ]
                    self.assertEqual(
                        [worksheet.cell(row=row, column=column).value for column in range(20, 45)],
                        values,
                    )
            finally:
                workbook.close()


if __name__ == "__main__":
    unittest.main()
