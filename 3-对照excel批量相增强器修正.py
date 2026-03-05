import cv2
import numpy as np
import tkinter as tk
from tkinter import filedialog
import os
from lib.lib_cv import *
import openpyxl

folder_path = r'D:\工作文件同步\20250826孔径比影响\处理图像\2-减底噪图'
correction_matrix_folder = r'D:\工作文件同步\20250826孔径比影响\处理图像\3.2-相增强器修正减底噪'
xlsx_path = r'D:\工作文件同步\20250826孔径比影响\20250820孔径比影响-正式-工况总结.xlsx'

if __name__ == "__main__":
    wb = openpyxl.load_workbook(xlsx_path)
    ws = wb.active

    img_paths = []
    for root, dirs, files in os.walk(folder_path):
        for file in files:
            if file.endswith('.npy'):
                full_path = os.path.join(root, file)
                img_paths.append(full_path)

    for img_path in img_paths:
        filename = os.path.basename(img_path)
        if '实验编号' not in filename:
            continue
        exp_num = int(os.path.splitext(filename)[0][4:])
        for row_num in range(1,ws.max_row+1):
            if ws.cell(row=row_num, column=1).value == exp_num:
                gain = int(ws.cell(row=row_num, column=4).value*100)
                correction_matrix = np.load(correction_matrix_folder + f'\\Gain{gain}.npy')
                process_image(img_path, correction_matrix, 'multiply')
                print(f'处理实验编号{exp_num}，增益{gain}')